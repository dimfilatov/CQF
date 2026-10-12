from io import BytesIO
from pathlib import Path
from typing import Iterable, Sequence

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Inches


SUMMARY_PATH = Path(__file__).with_name("data_engineering_summary.docx")

_ORIGINAL_FEATURES = {
    "open": "Opening price for the trading session.",
    "high": "Highest traded price during the trading session.",
    "low": "Lowest traded price during the trading session.",
    "close": "Closing price for the trading session.",
    "volume": "Number of shares traded during the trading session.",
}

_BASE_FEATURES = {
    "OC": ("Open-to-close price difference.", "Intraday directional price movement."),
    "HC": ("High-to-low price difference.", "Intraday trading range."),
    "GAP": ("Percentage gap up or down.", "Overnight repricing between sessions."),
    "RET": ("Daily log return of close.", "One-session proportional price movement."),
}

_WINDOW_FEATURES = {
    "PCHG": (
        "Close percentage change over the window.",
        "Price momentum over the specified lookback.",
    ),
    "VCHG": (
        "Volume percentage change over the window.",
        "Change in trading activity over the specified lookback.",
    ),
    "RET": (
        "Rolling sum of daily log returns over the window.",
        "Cumulative log price movement over the lookback.",
    ),
    "OC": (
        "Rolling mean of open-to-close differences.",
        "Average intraday directional movement over the lookback.",
    ),
    "HC": (
        "Rolling mean of high-to-low differences.",
        "Average intraday price range over the lookback.",
    ),
    "GAP": (
        "Rolling mean of percentage gaps.",
        "Average overnight repricing over the lookback.",
    ),
    "STD": (
        "Rolling standard deviation of daily log returns.",
        "Realized return volatility over the lookback.",
    ),
    "BB": (
        "Close position: (close - lower band) / (upper band - lower band); "
        "bands use two standard deviations.",
        "Relative price position within the Bollinger envelope; values can "
        "fall outside 0 to 1.",
    ),
    "RSI": (
        "Relative Strength Index calculated over the window.",
        "Momentum oscillator on a 0 to 100 scale.",
    ),
    "STOCH": (
        "Stochastic oscillator over the window with smooth_window=1.",
        "Close location within the recent high-low range, on a 0 to 100 scale.",
    ),
}


def _format_value(value: object) -> str:
    if value is None or value is pd.NA:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def _add_table(
    document: DocumentType,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = _format_value(value)


def _add_correlation_heatmap(document: DocumentType, corr: pd.DataFrame) -> None:
    """Render the same heatmap as plot_correlation_heatmap() into the report."""
    fig, ax = plt.subplots(figsize=(12, 10))
    im = ax.imshow(corr, cmap="RdBu_r", vmin=-1, vmax=1)
    ax.set_xticks(np.arange(len(corr.columns)))
    ax.set_yticks(np.arange(len(corr.index)))
    ax.set_xticklabels(corr.columns, rotation=90)
    ax.set_yticklabels(corr.index)
    for i in range(len(corr)):
        for j in range(len(corr.columns)):
            ax.text(j, i, f"{corr.iloc[i, j]:.2f}", ha="center", va="center", fontsize=8)
    fig.colorbar(im, ax=ax, label="Correlation")
    ax.set_title("Feature Correlation Matrix")
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=150)
    plt.close(fig)
    buffer.seek(0)
    document.add_picture(buffer, width=Inches(6.3))


def _add_dataframe(
    document: DocumentType, frame: pd.DataFrame, include_index: bool = False
) -> None:
    headers = [str(column) for column in frame.columns]
    if include_index:
        headers.insert(0, str(frame.index.name or ""))
    rows = (
        ((index, *values) if include_index else tuple(values))
        for index, values in zip(frame.index, frame.itertuples(index=False))
    )
    _add_table(document, headers, rows)


def _feature_rows(feature_columns: Sequence[str]) -> list[tuple[str, str, str]]:
    rows = []
    for feature in feature_columns:
        if feature in _BASE_FEATURES:
            transformation, interpretation = _BASE_FEATURES[feature]
        else:
            prefix = next(
                (
                    candidate
                    for candidate in _WINDOW_FEATURES
                    if feature.startswith(candidate)
                    and feature[len(candidate) :].isdigit()
                ),
                None,
            )
            if prefix is None:
                transformation = "Derived from the input price/volume series."
                interpretation = "Engineered model input."
            else:
                window = feature[len(prefix) :]
                transformation, interpretation = _WINDOW_FEATURES[prefix]
                transformation = f"{transformation} Window = {window}."
        rows.append((feature, transformation, interpretation))
    return rows


def Data_engineering_template(
    data_engineering, output_path: Path = SUMMARY_PATH
) -> Path:
    """Create the starter Word report and document the loaded input features."""
    document = Document()
    document.add_heading("Data Engineering Summary", 0)
    document.add_paragraph(
        "Generated by the exam_3 data pipeline. The report is extended as "
        "feature engineering and selection complete."
    )
    document.add_heading("Original features in self.data", level=1)
    original_features = data_engineering.original_feature_columns
    _add_table(
        document,
        ("Feature", "Meaning", "Data type", "Missing values after volume cleaning"),
        (
            (
                feature,
                _ORIGINAL_FEATURES.get(feature, "Original downloaded input feature."),
                data_engineering.data[feature].dtype,
                int(data_engineering.data[feature].isna().sum()),
            )
            for feature in original_features
        ),
    )

    document.add_heading("Loaded data and initial cleaning", level=1)
    first_date = data_engineering.data.index.min()
    last_date = data_engineering.data.index.max()
    _add_table(
        document,
        ("Item", "Result"),
        (
            ("Symbol", data_engineering.yahoo_symbol),
            ("Requested start date", data_engineering.start_date),
            ("Downloaded rows", data_engineering.downloaded_row_count),
            ("Rows removed for zero volume", data_engineering.rows_removed_for_zero_volume),
            ("Rows after zero-volume cleaning", len(data_engineering.data)),
            ("Date range after cleaning", f"{first_date} to {last_date}"),
        ),
    )
    document.add_paragraph(
        "This section is a starter summary. The following pipeline calls append "
        "the engineered-feature, row-cleaning, and selected-feature results."
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def data_engineering_word_summary(
    data_engineering, output_path: Path = SUMMARY_PATH
) -> Path:
    """Append transformations and pre-selection data cleaning results."""
    if not output_path.exists():
        raise FileNotFoundError(
            f"Word summary template does not exist: {output_path}. "
            "Call Data_engineering_template() after load_data() first."
        )
    document = Document(output_path)
    document.add_heading("New features created", level=1)
    _add_table(
        document,
        ("Feature", "Transformation", "Economic interpretation"),
        _feature_rows(data_engineering.engineered_feature_columns),
    )
    document.add_paragraph(
        "Target label: 1 when the next-session log return exceeds "
        f"{data_engineering.vol_scale} times the rolling 20-session standard "
        "deviation of daily log returns; otherwise 0. Rows without a valid "
        "forward return or threshold are removed."
    )

    document.add_heading("Dropped features and data cleaning", level=1)
    _add_table(
        document,
        ("Stage", "Action and result"),
        (
            (
                "Raw columns removed",
                "Dropped after deriving the technical features: "
                + ", ".join(data_engineering.original_feature_columns),
            ),
            (
                "Missing-value rows removed",
                f"{data_engineering.rows_removed_for_missing_values} of "
                f"{data_engineering.rows_before_missing_value_drop} rows were "
                "dropped because at least one engineered feature or the forward "
                "label was missing. This includes lookback warm-up and the final "
                "row without a forward return.",
            ),
            (
                "Missing-value treatment",
                "Rows with missing values were removed using dropna(); no "
                "imputation or explicit infinity replacement is performed.",
            ),
            (
                "Label treatment",
                "Rows without a valid forward-return label were excluded; "
                "remaining labels were converted to integers.",
            ),
            (
                "Feature-selection input",
                f"Chronological split: {len(data_engineering.X_train)} in-sample "
                f"rows and {len(data_engineering.X) - len(data_engineering.X_train)} "
                "out-of-sample rows. FeatureEngineering receives only the "
                "in-sample X and y.",
            ),
        ),
    )

    missing = data_engineering.missing_values_before_drop
    missing_rows = [
        (feature, int(count))
        for feature, count in missing.items()
        if count > 0
    ]
    document.add_paragraph("Missing values before row removal by column:")
    if missing_rows:
        _add_table(document, ("Column", "Missing values"), missing_rows)
    else:
        document.add_paragraph("No missing values were present before dropna().")

    x_train = data_engineering.X_train
    y_train = data_engineering.y_train
    document.add_heading("In-sample training data summary", level=1)
    first_date = x_train.index.min()
    last_date = x_train.index.max()
    _add_table(
        document,
        ("Item", "Result"),
        (
            ("X_train shape", f"{x_train.shape[0]} rows x {x_train.shape[1]} columns"),
            ("y_train length", len(y_train)),
            ("Date range", f"{first_date} to {last_date}"),
            ("Missing values in X_train", int(x_train.isna().sum().sum())),
            ("Missing values in y_train", int(y_train.isna().sum())),
        ),
    )
    document.add_paragraph("self.X_train descriptive statistics:")
    _add_dataframe(document, x_train.describe().T, include_index=True)
    document.add_paragraph("self.y_train descriptive statistics:")
    _add_dataframe(
        document, y_train.describe().to_frame(y_train.name or "y_train").T,
        include_index=True,
    )

    document.save(output_path)
    return output_path


def produce_label_imbalance_summary(
    data_engineering, output_path: Path = SUMMARY_PATH
) -> Path:
    """Append the self.label_imbalance table to the Word report."""
    if not output_path.exists():
        raise FileNotFoundError(
            f"Word summary document does not exist: {output_path}. "
            "Call Data_engineering_template() after load_data() first."
        )
    document = Document(output_path)
    document.add_heading("Label imbalance", level=1)
    document.add_paragraph(
        "Label counts and percentages for the in-sample and out-of-sample splits."
    )
    frame = data_engineering.label_imbalance.copy()
    frame["percentage"] = frame["percentage"].map("{:.2f}%".format)
    _add_dataframe(document, frame)
    document.save(output_path)
    return output_path


def feature_engineering_word_summary(
    data_engineering, output_path: Path = SUMMARY_PATH
) -> Path:
    """Append the selected candidates and their engineered-data summary."""
    if not output_path.exists():
        raise FileNotFoundError(
            f"Word summary document does not exist: {output_path}. "
            "Call Data_engineering_template() after load_data() first."
        )
    document = Document(output_path)
    selected_features = sorted(data_engineering.features)
    excluded_features = sorted(
        set(data_engineering.engineered_feature_columns) - set(selected_features)
    )
    document.add_heading("Feature selection results", level=1)
    document.add_paragraph(
        "Selected candidates are the union of features returned by the "
        "configured selection algorithms."
    )
    _add_table(
        document,
        ("Selection result", "Features"),
        (
            ("Selected candidates", ", ".join(selected_features)),
            ("Number selected", len(selected_features)),
            ("Engineered candidates not selected", ", ".join(excluded_features)),
            ("Number not selected", len(excluded_features)),
            (
                "Selection algorithms",
                ", ".join(data_engineering.feature_selection_methods),
            ),
        ),
    )

    document.add_heading("Selected features (self.features)", level=1)
    _add_table(
        document,
        ("#", "Feature"),
        ((i, feature) for i, feature in enumerate(data_engineering.features, 1)),
    )

    document.add_heading("Feature correlation heatmap (self.corr)", level=1)
    _add_correlation_heatmap(document, data_engineering.corr)

    document.add_heading("Selection by algorithm (self.feature_df)", level=1)
    _add_dataframe(document, data_engineering.feature_df)

    document.add_heading("summarize_engineered_data() results", level=1)
    summary = data_engineering.engineered_data_summary
    _add_table(
        document,
        ("Feature", *[str(column) for column in summary.columns]),
        (
            (feature, *summary.loc[feature].tolist())
            for feature in summary.index
        ),
    )
    document.save(output_path)
    return output_path
